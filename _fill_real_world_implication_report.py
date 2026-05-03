from pathlib import Path

from docx import Document


TEMPLATE = Path(r"c:\Saadat\Class\Spring 2026\CSE461\Project\Real World Implication Report Template.docx")
OUTPUT = Path(r"c:\Saadat\Class\Spring 2026\CSE461\Project\Real World Implication Report_FILLED_Group5.docx")


def fill_report() -> None:
    doc = Document(str(TEMPLATE))

    # Header metadata
    paragraph_updates = {
        10: "Group : 5",
        11: "Section : 3",
        12: "Semester : Spring_2026",
        19: "Submitted Date: 26-04-2026",
        20: "Submitted to - RDW, SMYA",
    }

    for idx, text in paragraph_updates.items():
        if idx < len(doc.paragraphs):
            doc.paragraphs[idx].text = text

    # Body content aligned with template section placeholders.
    section_paragraph_updates = {
        22: (
            "The Smart Kitchen Safety and Ventilation System is an autonomous stationary safety platform "
            "designed for home and shared kitchens. It continuously observes gas leakage trend, temperature "
            "and humidity behavior, and flame presence, then uses a multi-condition state machine to decide "
            "whether the environment is normal, warning-level, or dangerous."
        ),
        25: (
            "The purpose of this project is to reduce kitchen fire and gas-related accidents by combining early "
            "detection with automatic response. Instead of depending only on manual reaction, the system can "
            "activate alarm, ventilation, and gas-cutoff protection while also reducing false alarms during "
            "normal cooking."
        ),
        28: (
            "This project addresses a common real-world problem in Bangladesh where LPG-based cooking is widely "
            "used in homes, hostels, and rental apartments. A low-cost intelligent safety system can improve "
            "household security, protect children and elderly family members, and increase public awareness of "
            "preventive kitchen safety practices."
        ),
        31: (
            "By detecting dangerous gas and heat conditions early, the system can lower exposure to harmful smoke "
            "and leaked gas, potentially reducing respiratory stress and burn injuries. Better kitchen ventilation "
            "control also supports healthier indoor air quality during long cooking sessions."
        ),
        34: (
            "Safety impact is the strongest implication of this project because it introduces automatic protective "
            "actions, not only notification. When risk persists, the system can trigger buzzer and visual alerts, "
            "run the fan, open vent flap, and lock gas cutoff. The Stop Alarm button is designed to mute sound "
            "without disabling protection, which prevents unsafe user override during emergencies."
        ),
        37: (
            "For practical deployment, electrical wiring and relay switching must follow local safety regulations "
            "and component ratings. Real installations should be tested and verified by qualified personnel, "
            "because failures in power design or safety-control logic could create liability issues. "
            "The project does not collect personal multimedia data, so privacy risk is low."
        ),
        40: (
            "Cooking behavior in Bangladeshi families is diverse and frequent, so the system must respect normal "
            "kitchen activity without creating frequent nuisance alarms. Clear status messages and simple controls "
            "make the design more acceptable for users of different ages and technical backgrounds, supporting a "
            "culture of preventive safety at home."
        ),
        43: (
            "The Smart Kitchen Safety and Ventilation System has strong positive real-world implications in social, "
            "health, and safety dimensions. Its value comes from combining multi-sensor intelligence with automatic "
            "action, making it more practical than single-threshold alarm devices. With proper testing and compliant "
            "installation, it can become a reliable low-cost model for safer kitchens."
        ),
    }

    for idx, text in section_paragraph_updates.items():
        if idx < len(doc.paragraphs):
            doc.paragraphs[idx].text = text

    # Fill the 5-row member table in the template.
    members = [
        ("Md.Abir Hasan Rohan", "22101294"),
        ("Md.Saadat Rahman", "22201101"),
        ("Tazwar Saif Chowdhury", "22201373"),
        ("Muntasir Mamun Sakib", "22299146"),
        ("Rufaida Faruque", "23101005"),
    ]

    if doc.tables:
        table = doc.tables[0]
        for row_idx, (name, student_id) in enumerate(members):
            if row_idx < len(table.rows) and len(table.rows[row_idx].cells) >= 2:
                table.cell(row_idx, 0).text = name
                table.cell(row_idx, 1).text = student_id

    doc.save(str(OUTPUT))
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    fill_report()
