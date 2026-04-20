from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt
from pathlib import Path

OUT = Path(r"c:\Saadat\Class\Spring 2026\CSE461\Project\CSE461 Project Proposal Template_FILLED_Group5.docx")

ideas = [
    {
        "title": "Smart Kitchen Safety and Ventilation System",
        "objective": "Design and build an autonomous kitchen safety system that uses multi-condition logic to distinguish normal cooking from dangerous events, reduce false alarms, and automatically activate ventilation, alarms, and gas-cutoff safety actions.",
        "scope": "The project is intended for household kitchens, dorm kitchens, or small food preparation areas. The system continuously monitors gas trend, temperature-humidity condition, and flame presence near the cooking area (with optional ambient-light context). Based on a multi-condition risk score and state-machine logic, it decides whether to open a vent, run an exhaust fan, trigger alarms, and execute gas cutoff.",
        "significance": "LPG gas leakage is a real and dangerous household problem in Bangladesh; most low-cost systems only raise an alarm, while this idea adds automatic action; it is socially useful and faculty-relevant; online implementation resources are widely available.",
        "components": [
            ("Microcontroller", "Arduino Uno"),
            ("Sensors", "MQ-2, DHT22, flame sensor, optional LDR"),
            ("Actuators", "SG90 servo, fan, solenoid gas valve, I2C LCD"),
            ("Body/Chassis", "DIY enclosure/frame"),
            ("Others", "Stop button, LEDs, buzzer, relay module(s), distribution board, breadboard, wires, power supply"),
        ],
        "cost_rows": [
            ("Arduino", "1x Uno", "700"),
            ("MQ-2", "1x sensor", "200"),
            ("DHT22", "1x sensor", "300"),
            ("Flame Sensor", "1x module", "120"),
            ("LDR set", "Optional LDR + resistors", "30"),
            ("Servo", "1x SG90", "200"),
            ("Fan", "1x exhaust/DC fan", "100"),
            ("Solenoid Valve", "1x gas valve", "550"),
            ("Relay", "1x 2-channel or equivalent", "150"),
            ("LCD", "1x 16x2 I2C", "350"),
            ("LED/Buzzer", "LEDs + buzzer", "80"),
            ("Stop Button", "1x push button", "20"),
            ("Board/Wires", "Breadboard + jumpers", "200"),
            ("Frame", "DIY frame", "150"),
            ("Supply", "Power supply set", "250"),
            ("Distribution Board", "1x fused terminal board", "150"),
            ("Total", "", "3550"),
        ],
    },
    {
        "title": "Smart Classroom Occupancy and Environment Controller",
        "objective": "Build an autonomous room monitoring and control system that counts people entering and leaving a room, monitors environmental conditions, and automatically controls ventilation and lighting to reduce energy waste.",
        "scope": "The project is designed for classrooms, offices, labs, or meeting rooms. It detects entry/exit movement, tracks occupancy count, and adjusts fan, vent, and lighting according to room usage and environmental readings.",
        "significance": "This solves visible real-world energy waste in empty rooms, is directly relevant to academic buildings, demonstrates multi-sensor automation logic, and has strong online support.",
        "components": [
            ("Microcontroller", "Arduino Uno"),
            ("Sensors", "2x IR, DHT11, LDR"),
            ("Actuators", "5V fan, SG90 servo, I2C LCD"),
            ("Body/Chassis", "DIY frame/enclosure"),
            ("Others", "LEDs, buzzer, button, relay, breadboard, wires, 5V supply"),
        ],
        "cost_rows": [
            ("Arduino", "1x Uno", "700"),
            ("IR pair", "2x FC-51", "90"),
            ("DHT11", "1x sensor", "120"),
            ("LDR set", "LDR + resistor", "16"),
            ("Fan", "1x 5V mini fan", "190"),
            ("Servo", "1x SG90", "150"),
            ("Relay", "1x 5V module", "81"),
            ("LCD", "1x 16x2 I2C", "155"),
            ("LED/Buzz/Button", "LEDs + buzzer + button", "29"),
            ("Board/Wires", "Breadboard + jumpers", "190"),
            ("Frame", "DIY frame", "200"),
            ("Supply", "5V 2A adapter", "150"),
            ("Total", "", "2071"),
        ],
    },
    {
        "title": "Smart Automated Parking Gate System with Vehicle Counting",
        "objective": "Design and build an autonomous parking gate system that detects vehicles, opens and closes a gate automatically, counts available parking spaces, and denies entry when the lot is full.",
        "scope": "This project is intended for small apartment, office, or campus parking areas. It simulates an automated gate controller using entry/exit detection, servo barrier control, and a live parking count display.",
        "significance": "Parking management is a practical urban problem; this project demonstrates clear automation with strong demo value and abundant online references.",
        "components": [
            ("Microcontroller", "Arduino Uno"),
            ("Sensors", "HC-SR04, IR, LDR"),
            ("Actuators", "SG90 servo, I2C LCD, LEDs, buzzer"),
            ("Body/Chassis", "Small parking model/frame"),
            ("Others", "Button, breadboard, wires, resistors, 5V supply"),
        ],
        "cost_rows": [
            ("Arduino", "1x Uno", "700"),
            ("HC-SR04", "1x sensor", "150"),
            ("IR", "1x FC-51", "45"),
            ("LDR set", "LDR + resistor", "16"),
            ("Servo", "1x SG90", "150"),
            ("LCD", "1x 16x2 I2C", "155"),
            ("LED/Buzzer", "LEDs + buzzer", "24"),
            ("Button", "1x push button", "5"),
            ("Board/Wires", "Breadboard + jumpers", "190"),
            ("Frame", "DIY parking model", "200"),
            ("Supply", "5V 2A adapter", "150"),
            ("Total", "", "1785"),
        ],
    },
]


def set_doc_defaults(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)


def compact_paragraph(paragraph, size=10, bold=False):
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    for run in paragraph.runs:
        run.font.size = Pt(size)
        run.font.bold = bold


def compact_table(table, col_widths, header_size=9, body_size=9):
    table.style = "Table Grid"
    table.autofit = True
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.width = col_widths[index]
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1
                for run in paragraph.runs:
                    run.font.size = Pt(body_size)
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(header_size)
                run.font.bold = True

new_doc = Document()
set_doc_defaults(new_doc)

header_lines = [
    "CSE461: Introduction to Robotics Lab",
    "Lab No: 2",
    "Section : 3",
    "Group : 5",
    "Semester : Spring_2026",
    "",
    "Group Members:",
    "Name: Md.Abir Hasan Rohan    ID: 22101294",
    "Name: Md.Saadat Rahman       ID: 22201101",
    "Name: Tazwar Saif Chowdhury  ID: 22201373",
    "Name: Muntasir Mamun Sakib   ID: 22299146",
    "Name: Rufaida Faruque        ID: 23101005",
    "",
    "Submitted to:",
    "Instructors: RDW, SMYA",
    "",
    "Project Proposal",
]

for line in header_lines:
    paragraph = new_doc.add_paragraph(line)
    compact_paragraph(paragraph, size=10, bold=(line in {"Group Members:", "Submitted to:", "Project Proposal"}))

for idea_idx, idea in enumerate(ideas, start=1):
    compact_paragraph(new_doc.add_paragraph(f"Project Idea {idea_idx}"), size=10, bold=True)
    compact_paragraph(new_doc.add_paragraph("1. Project Title:"), size=10, bold=True)
    compact_paragraph(new_doc.add_paragraph(idea["title"]), size=10)
    compact_paragraph(new_doc.add_paragraph("2. Purpose"), size=10, bold=True)
    compact_paragraph(new_doc.add_paragraph(f"Objective: {idea['objective']}"), size=10)
    compact_paragraph(new_doc.add_paragraph(f"Scope: {idea['scope']}"), size=10)
    compact_paragraph(new_doc.add_paragraph(f"Significance: {idea['significance']}"), size=10)
    compact_paragraph(new_doc.add_paragraph("3. Components"), size=10, bold=True)
    t1 = new_doc.add_table(rows=1, cols=2)
    t1.rows[0].cells[0].text = 'Type'
    t1.rows[0].cells[1].text = 'Items'
    for a, b in idea['components']:
        row = t1.add_row().cells
        row[0].text = a
        row[1].text = b
    compact_table(t1, [Inches(1.55), Inches(5.0)])
    compact_paragraph(new_doc.add_paragraph("4. Cost Breakdown (Approx.)"), size=10, bold=True)
    t2 = new_doc.add_table(rows=1, cols=3)
    t2.rows[0].cells[0].text = 'Item'
    t2.rows[0].cells[1].text = 'Spec'
    t2.rows[0].cells[2].text = 'BDT'
    for a, b, c in idea['cost_rows']:
        row = t2.add_row().cells
        row[0].text = a
        row[1].text = b
        row[2].text = c
    compact_table(t2, [Inches(1.6), Inches(4.4), Inches(0.9)])
    compact_paragraph(new_doc.add_paragraph("5. Functionality Breakdown"), size=10, bold=True)
    if idea_idx == 1:
        funcs = [
            ("Multi-Condition Hazard Detection with False Alarm Prevention", "The system uses multiple sensors and state-based logic so normal cooking activity does not immediately trigger danger actions.", "1) Arduino reads MQ-2, DHT22, and flame-sensor values continuously (optional LDR context can be added). 2) Sensor data is stabilized using warm-up lockout, moving-average filtering, and outlier rejection. 3) A weighted risk score determines states such as NORMAL, COOKING_NORMAL, WARNING, and ALARM_ACTIVE. 4) Persistence timers and hysteresis prevent one-sample spikes from creating false alarms. 5) Only sustained multi-condition danger transitions the system toward cutoff-level response."),
            ("Automatic Cutoff and Safety Actuation", "When confirmed danger persists, the system performs automatic protective actions including gas cutoff.", "1) If risk remains high for the configured confirmation window, the system enters ALARM_ACTIVE and then CUTOFF_LOCKED. 2) The solenoid gas valve is closed immediately via relay/driver. 3) The exhaust fan is turned on and vent flap is opened using the servo motor. 4) Buzzer, LEDs, and LCD provide clear alarm and hazard status. 5) Gas restore is blocked until a continuous safe window is observed and manual reset conditions are met."),
            ("Stop Alarm Button and Safe Recovery Workflow", "The Stop Alarm button improves usability without allowing users to bypass safety protection.", "1) A short press of the Stop Alarm button mutes the buzzer for a temporary window. 2) During mute, safety actions remain active: fan on, vent open, and gas cutoff locked. 3) LCD shows safety-active status (for example, ALARM MUTED - SAFETY ACTIVE). 4) A long press is accepted only after safe readings persist for the configured recovery window. 5) On valid recovery, the system returns to normal state and allows controlled reset of alarm status."),
        ]
    elif idea_idx == 2:
        funcs = [
            ("Bidirectional Occupancy Counting", "Two IR sensors identify whether a person is entering or leaving the room based on trigger order.", "1) Two IR sensors are mounted at the doorway. 2) Outer-then-inner sequence counts entry. 3) Inner-then-outer sequence counts exit. 4) LCD shows live count and capacity. 5) If capacity is exceeded, buzzer and warning are shown."),
            ("Temperature-Controlled Ventilation", "When occupied, fan and vent are controlled according to temperature and humidity.", "1) DHT11 monitors temperature/humidity continuously. 2) Empty room keeps fan off and vent closed. 3) High temperature while occupied turns on fan via relay. 4) Servo opens vent partially/fully by threshold. 5) LCD shows room and ventilation status."),
            ("Occupancy-Based Smart Lighting", "LDR controls room lights based on ambient light only when the room is occupied.", "1) LDR provides analog light level. 2) Empty room keeps lights off. 3) Occupied + low light turns on white LEDs. 4) Medium light enables partial lighting. 5) LCD alternates occupancy, environment, and lighting status."),
        ]
    else:
        funcs = [
            ("Vehicle Detection and Gate Opening", "The system detects an incoming vehicle and opens the gate automatically if space is available.", "1) Ultrasonic sensor detects vehicle at entry lane. 2) Arduino checks parking count. 3) If not full, servo raises barrier. 4) After passing, gate closes. 5) LCD shows gate and lot status."),
            ("Exit Detection and Space Count Update", "An IR sensor detects vehicles leaving and updates available space count.", "1) IR sensor detects passing vehicle at exit. 2) Arduino decrements occupied count. 3) LCD updates available spots. 4) Full lot returns to entry-allowed once count drops. 5) Push button supports manual reset during testing."),
            ("Day/Night Status Lighting and Alert Indication", "LDR changes lighting and status behavior between day and night conditions.", "1) LDR reads ambient brightness. 2) Day mode uses standard indicators. 3) Night mode enables additional white LEDs. 4) Full lot triggers red LED and buzzer. 5) LCD alternates parking count, gate state, and day/night mode."),
        ]
    for idx, (n,o,w) in enumerate(funcs, start=1):
        compact_paragraph(new_doc.add_paragraph(f"Functionality {idx}: {n}"), size=10, bold=True)
        compact_paragraph(new_doc.add_paragraph(f"Overview: {o}"), size=10)
        compact_paragraph(new_doc.add_paragraph(f"Working Procedure: {w}"), size=10)

# Save as separate filled file with tables
new_doc.save(str(OUT))
print(f"Saved filled file with tables: {OUT}")
