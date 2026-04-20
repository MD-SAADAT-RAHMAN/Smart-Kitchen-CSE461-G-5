import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from docx import Document

doc = Document(r"Project\Hardware Project Guideline.docx")

print("=" * 60)
print("ALL PARAGRAPHS")
print("=" * 60)
for i, p in enumerate(doc.paragraphs, 1):
    print(f"[{i}] Style={p.style.name} | {repr(p.text)}")

print()
print("=" * 60)
print("ALL TABLES")
print("=" * 60)
for t_idx, table in enumerate(doc.tables, 1):
    print(f"\n--- TABLE {t_idx} ({len(table.rows)} rows x {len(table.columns)} cols) ---")
    for r_idx, row in enumerate(table.rows, 1):
        for c_idx, cell in enumerate(row.cells, 1):
            ct = cell.text.strip()
            if ct:
                print(f"  R{r_idx}C{c_idx}: {ct}")
