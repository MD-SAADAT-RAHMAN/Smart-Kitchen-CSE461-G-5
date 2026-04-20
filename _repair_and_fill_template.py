from docx import Document
from pathlib import Path

p = Path(r"c:\Saadat\Class\Spring 2026\CSE461\Project\CSE461 Project Proposal Template.docx")
doc = Document(str(p))

# Header and member details
header_updates = {
    2: "CSE461",
    3: "Section : 3",
    4: "Group : 5",
    5: "Semester : Spring_2026",
    6: "Robot Project Proposal",
    8: "Member Details:",
    9: "Name: Md.Abir Hasan Rohan    ID: 22101294",
    10: "Name: Md.Saadat Rahman       ID: 22201101",
    11: "Name: Tazwar Saif Chowdhury  ID: 22201373",
    12: "Name: Muntasir Mamun Sakib   ID: 22299146",
    13: "Name: Rufaida Faruque        ID: 23101005",
    15: "Submission Date:",
    16: "Instructors: RDW, SMYA",
}
for i, v in header_updates.items():
    if i < len(doc.paragraphs):
        doc.paragraphs[i].text = v

ideas = [
    {
        "title": "Smart Kitchen Safety and Ventilation System",
        "objective": "Design and build an autonomous kitchen safety system that uses multi-condition logic to distinguish normal cooking from dangerous events, reduce false alarms, and automatically activate ventilation, alarms, and gas-cutoff safety actions.",
        "scope": "The project is intended for household kitchens, dorm kitchens, or small food preparation areas. The system continuously monitors gas trend, temperature-humidity condition, and flame presence near the cooking area (with optional ambient-light context). Based on a multi-condition risk score and state-machine logic, it decides whether to open a vent, run an exhaust fan, trigger alarms, and execute gas cutoff.",
        "significance": "LPG gas leakage is a real and dangerous household problem in Bangladesh; most low-cost systems only raise an alarm, while this idea adds automatic action; it is socially useful and faculty-relevant; online implementation resources are widely available.",
        "micro": "Arduino Uno",
        "sensors": "MQ-2 gas/smoke sensor; DHT22 temperature and humidity sensor; Flame sensor module (IR); Optional LDR for kitchen activity / ambient-light context",
        "actuators": "SG90 servo motor for ventilation flap; Exhaust/DC fan controlled by relay; Solenoid gas valve for automatic cutoff (relay or MOSFET driver); I2C LCD 16x2",
        "body": "DIY enclosure/frame",
        "additional": "Stop Alarm push button (mandatory), LEDs, buzzer, relay module(s), power distribution board (recommended), resistors, breadboard, jumper wires, power supply",
        "cost": "Total (Approx.): 3,550 BDT",
        "f1n": "Multi-Condition Hazard Detection with False Alarm Prevention",
        "f1o": "The system uses multiple sensors and state-based logic so normal cooking activity does not immediately trigger danger actions.",
        "f1w": "1) Arduino reads MQ-2, DHT22, and flame-sensor values continuously (optional LDR context can be added). 2) Sensor data is stabilized using warm-up lockout, moving-average filtering, and outlier rejection. 3) A weighted risk score determines states such as NORMAL, COOKING_NORMAL, WARNING, and ALARM_ACTIVE. 4) Persistence timers and hysteresis prevent one-sample spikes from creating false alarms. 5) Only sustained multi-condition danger transitions the system toward cutoff-level response.",
        "f2n": "Automatic Cutoff and Safety Actuation",
        "f2o": "When confirmed danger persists, the system performs automatic protective actions including gas cutoff.",
        "f2w": "1) If risk remains high for the configured confirmation window, the system enters ALARM_ACTIVE and then CUTOFF_LOCKED. 2) The solenoid gas valve is closed immediately via relay/driver. 3) The exhaust fan is turned on and vent flap is opened using the servo motor. 4) Buzzer, LEDs, and LCD provide clear alarm and hazard status. 5) Gas restore is blocked until a continuous safe window is observed and manual reset conditions are met.",
        "f3n": "Stop Alarm Button and Safe Recovery Workflow",
        "f3o": "The Stop Alarm button improves usability without allowing users to bypass safety protection.",
        "f3w": "1) A short press of the Stop Alarm button mutes the buzzer for a temporary window. 2) During mute, safety actions remain active: fan on, vent open, and gas cutoff locked. 3) LCD shows safety-active status (for example, ALARM MUTED - SAFETY ACTIVE). 4) A long press is accepted only after safe readings persist for the configured recovery window. 5) On valid recovery, the system returns to normal state and allows controlled reset of alarm status.",
    },
    {
        "title": "Smart Classroom Occupancy and Environment Controller",
        "objective": "Build an autonomous room monitoring and control system that counts people entering and leaving a room, monitors environmental conditions, and automatically controls ventilation and lighting to reduce energy waste.",
        "scope": "The project is designed for classrooms, offices, labs, or meeting rooms. It detects entry/exit movement, tracks occupancy count, and adjusts fan, vent, and lighting according to room usage and environmental readings.",
        "significance": "This solves visible real-world energy waste in empty rooms, is directly relevant to academic buildings, demonstrates multi-sensor automation logic, and has strong online support.",
        "micro": "Arduino Uno",
        "sensors": "IR sensor pair for entry/exit counting; DHT11 temperature and humidity sensor; LDR for ambient light detection",
        "actuators": "DC fan controlled by relay; SG90 servo motor for vent/window flap; I2C LCD 16x2",
        "body": "DIY frame/enclosure",
        "additional": "White LEDs for room light simulation, buzzer, push button, relay module, breadboard, jumper wires, power supply",
        "cost": "Total (Approx.): 2,071 BDT",
        "f1n": "Bidirectional Occupancy Counting",
        "f1o": "Two IR sensors identify whether a person is entering or leaving the room based on trigger order.",
        "f1w": "1) Two IR sensors are mounted at the doorway. 2) Outer-then-inner sequence counts entry. 3) Inner-then-outer sequence counts exit. 4) LCD shows live count and capacity. 5) If capacity is exceeded, buzzer and warning are shown.",
        "f2n": "Temperature-Controlled Ventilation",
        "f2o": "When occupied, fan and vent are controlled according to temperature and humidity.",
        "f2w": "1) DHT11 monitors temperature/humidity continuously. 2) Empty room keeps fan off and vent closed. 3) High temperature while occupied turns on fan via relay. 4) Servo opens vent partially/fully by threshold. 5) LCD shows room and ventilation status.",
        "f3n": "Occupancy-Based Smart Lighting",
        "f3o": "LDR controls room lights based on ambient light only when the room is occupied.",
        "f3w": "1) LDR provides analog light level. 2) Empty room keeps lights off. 3) Occupied + low light turns on white LEDs. 4) Medium light enables partial lighting. 5) LCD alternates occupancy, environment, and lighting status.",
    },
    {
        "title": "Smart Automated Parking Gate System with Vehicle Counting",
        "objective": "Design and build an autonomous parking gate system that detects vehicles, opens and closes a gate automatically, counts available parking spaces, and denies entry when the lot is full.",
        "scope": "This project is intended for small apartment, office, or campus parking areas. It simulates an automated gate controller using entry/exit detection, servo barrier control, and a live parking count display.",
        "significance": "Parking management is a practical urban problem; this project demonstrates clear automation with strong demo value and abundant online references.",
        "micro": "Arduino Uno",
        "sensors": "HC-SR04 ultrasonic sensor for entry detection; IR sensor for exit detection; LDR for day/night mode",
        "actuators": "SG90 servo motor for barrier gate; I2C LCD 16x2; LEDs and buzzer for traffic/status indication",
        "body": "Small parking model/frame",
        "additional": "Push button, breadboard, jumper wires, resistors, power supply",
        "cost": "Total (Approx.): 1,785 BDT",
        "f1n": "Vehicle Detection and Gate Opening",
        "f1o": "The system detects an incoming vehicle and opens the gate automatically if space is available.",
        "f1w": "1) Ultrasonic sensor detects vehicle at entry lane. 2) Arduino checks parking count. 3) If not full, servo raises barrier. 4) After passing, gate closes. 5) LCD shows gate and lot status.",
        "f2n": "Exit Detection and Space Count Update",
        "f2o": "An IR sensor detects vehicles leaving and updates available space count.",
        "f2w": "1) IR sensor detects passing vehicle at exit. 2) Arduino decrements occupied count. 3) LCD updates available spots. 4) Full lot returns to entry-allowed once count drops. 5) Push button supports manual reset during testing.",
        "f3n": "Day/Night Status Lighting and Alert Indication",
        "f3o": "LDR changes lighting and status behavior between day and night conditions.",
        "f3w": "1) LDR reads ambient brightness. 2) Day mode uses standard indicators. 3) Night mode enables additional white LEDs. 4) Full lot triggers red LED and buzzer. 5) LCD alternates parking count, gate state, and day/night mode.",
    },
]

# Fixed index map (from current 94-paragraph template)
starts = [19, 45, 70]
for s, idea in zip(starts, ideas):
    mapping = {
        s + 0: f"Project Idea {starts.index(s)+1}",
        s + 1: "1. Project Title:",
        s + 2: idea["title"],
        s + 3: "2. Purpose",
        s + 4: f"Objective: {idea['objective']}",
        s + 5: f"Scope: {idea['scope']}",
        s + 6: f"Significance: {idea['significance']}",
        s + 7: "3. Components",
        s + 8: f"Microcontroller: {idea['micro']}",
        s + 9: f"Sensors: {idea['sensors']}",
        s + 10: f"Actuators: {idea['actuators']}",
        s + 11: f"Body/Chassis: {idea['body']}",
        s + 12: f"Additional Components: {idea['additional']}",
        s + 13: f"4. Cost Breakdown (Approx.): {idea['cost']}",
        s + 14: "5. Functionality Breakdown",
        s + 15: f"Functionality 1: {idea['f1n']}",
        s + 16: f"Overview: {idea['f1o']}",
        s + 17: f"Working Procedure: {idea['f1w']}",
        s + 18: f"Functionality 2: {idea['f2n']}",
        s + 19: f"Overview: {idea['f2o']}",
        s + 20: f"Working Procedure: {idea['f2w']}",
        s + 21: f"Functionality 3: {idea['f3n']}",
        s + 22: f"Overview: {idea['f3o']}",
        s + 23: f"Working Procedure: {idea['f3w']}",
    }
    # keep separator/blank after each block
    if s == 19:
        mapping[s + 24] = ""
    elif s == 45:
        mapping[s + 24] = ""

    for idx, text in mapping.items():
        if idx < len(doc.paragraphs):
            doc.paragraphs[idx].text = text

# Save

doc.save(str(p))
print(f"Repaired and filled: {p}")
