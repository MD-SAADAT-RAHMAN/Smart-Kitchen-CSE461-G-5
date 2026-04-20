from docx import Document
from pathlib import Path

TEMPLATE = Path(r"c:\Saadat\Class\Spring 2026\CSE461\Project\CSE461 Project Proposal Template.docx")
OUT = Path(r"c:\Saadat\Class\Spring 2026\CSE461\Project\CSE461 Project Proposal Template.docx")

# First 3 ideas from Final Project Ideas.docx
ideas = [
    {
        "title": "Smart Kitchen Safety and Ventilation System",
        "objective": "Design and build an autonomous kitchen safety system that uses multi-condition logic to distinguish normal cooking from dangerous events, reduce false alarms, and automatically activate ventilation, alarms, and gas-cutoff safety actions.",
        "scope": "The project is intended for household kitchens, dorm kitchens, or small food preparation areas. The system continuously monitors gas trend, temperature-humidity condition, and flame presence near the cooking area (with optional ambient-light context). Based on a multi-condition risk score and state-machine logic, it decides whether to open a vent, run an exhaust fan, trigger alarms, and execute gas cutoff.",
        "significance": "1) LPG gas leakage is a real and dangerous household problem in Bangladesh. 2) Most low-cost systems only raise an alarm; this idea goes further by taking automatic action. 3) It is highly relevant, socially useful, and likely to be appreciated by faculty. 4) Similar systems are widely available online, so implementation help and troubleshooting resources are easy to find.",
        "micro": "Arduino Uno",
        "sensors": "MQ-2 gas/smoke sensor; DHT22 temperature and humidity sensor; Flame sensor module (IR); Optional LDR for kitchen activity / ambient-light context",
        "actuators": "SG90 servo motor for ventilation flap; Exhaust/DC fan controlled by relay; Solenoid gas valve for automatic cutoff (relay or MOSFET driver); I2C LCD 16x2",
        "body": "DIY enclosure/frame",
        "additional": "Stop Alarm push button (mandatory), LEDs, buzzer, relay module(s), power distribution board (recommended), resistors, breadboard, jumper wires, power supply",
        "cost": "Total (Approx.): 3,550 BDT",
        "f1": ("Multi-Condition Hazard Detection with False Alarm Prevention", "The system uses multiple sensors and state-based logic so normal cooking activity does not immediately trigger danger actions.", "1) Arduino reads MQ-2, DHT22, and flame-sensor values continuously (optional LDR context can be added). 2) Sensor data is stabilized using warm-up lockout, moving-average filtering, and outlier rejection. 3) A weighted risk score determines states such as NORMAL, COOKING_NORMAL, WARNING, and ALARM_ACTIVE. 4) Persistence timers and hysteresis prevent one-sample spikes from creating false alarms. 5) Only sustained multi-condition danger transitions the system toward cutoff-level response."),
        "f2": ("Automatic Cutoff and Safety Actuation", "When confirmed danger persists, the system performs automatic protective actions including gas cutoff.", "1) If risk remains high for the configured confirmation window, the system enters ALARM_ACTIVE and then CUTOFF_LOCKED. 2) The solenoid gas valve is closed immediately via relay/driver. 3) The exhaust fan is turned on and vent flap is opened using the servo motor. 4) Buzzer, LEDs, and LCD provide clear alarm and hazard status. 5) Gas restore is blocked until a continuous safe window is observed and manual reset conditions are met."),
        "f3": ("Stop Alarm Button and Safe Recovery Workflow", "The Stop Alarm button improves usability without allowing users to bypass safety protection.", "1) A short press of the Stop Alarm button mutes the buzzer for a temporary window. 2) During mute, safety actions remain active: fan on, vent open, and gas cutoff locked. 3) LCD shows safety-active status (for example, ALARM MUTED - SAFETY ACTIVE). 4) A long press is accepted only after safe readings persist for the configured recovery window. 5) On valid recovery, the system returns to normal state and allows controlled reset of alarm status."),
    },
    {
        "title": "Smart Classroom Occupancy and Environment Controller",
        "objective": "Build an autonomous room monitoring and control system that counts people entering and leaving a room, monitors environmental conditions, and automatically controls ventilation and lighting to reduce energy waste.",
        "scope": "The project is designed for classrooms, offices, labs, or meeting rooms. It detects entry/exit movement, tracks occupancy count, and adjusts fan, vent, and lighting conditions according to room usage and environmental readings.",
        "significance": "1) This solves a very visible real-world problem: fans and lights often remain on in empty rooms. 2) It is directly relevant to academic buildings, so faculty can immediately relate to its usefulness. 3) It demonstrates logic, automation, and multi-sensor decision-making rather than simple sensor reading. 4) There are many Arduino tutorials for people counters, smart room automation, and occupancy-based energy systems.",
        "micro": "Arduino Uno",
        "sensors": "IR sensor pair for entry/exit counting; DHT11 temperature and humidity sensor; LDR for ambient light detection",
        "actuators": "DC fan controlled by relay; SG90 servo motor for vent/window flap; I2C LCD 16x2",
        "body": "DIY frame/enclosure",
        "additional": "White LEDs for room light simulation, buzzer, push button, relay module, breadboard, jumper wires, power supply",
        "cost": "Total (Approx.): 2,071 BDT",
        "f1": ("Bidirectional Occupancy Counting", "Two IR sensors identify whether a person is entering or leaving the room based on trigger order. The occupancy count is updated automatically.", "1) Two IR sensors are placed at the doorway with a small gap between them. 2) If the outer sensor triggers before the inner sensor, Arduino counts it as entry. 3) If the inner sensor triggers before the outer sensor, Arduino counts it as exit. 4) The LCD shows the live room count and maximum capacity. 5) If occupancy reaches a limit, the buzzer sounds briefly and the LCD displays a full-room warning."),
        "f2": ("Temperature-Controlled Ventilation", "When the room is occupied, the fan and vent are controlled according to temperature and humidity readings.", "1) The DHT11 continuously measures temperature and humidity. 2) If the room is empty, the system keeps the fan off and closes the vent to save power. 3) If the room is occupied and temperature rises above the comfort threshold, the relay turns the fan on. 4) The servo motor opens the vent partially or fully depending on temperature level. 5) The LCD displays messages such as ROOM EMPTY, FAN ON, or HIGH TEMP."),
        "f3": ("Occupancy-Based Smart Lighting", "The LDR measures ambient light, and the system turns on LED room lights only when the room is occupied and natural light is insufficient.", "1) The LDR provides an analog light-level reading. 2) If the room is empty, all room lights remain off regardless of light level. 3) If the room is occupied and the light level is low, Arduino turns on white LEDs. 4) If the room is occupied and the light level is medium, only part of the LEDs turn on. 5) The LCD alternates between occupancy status, environmental readings, and lighting status."),
    },
    {
        "title": "Smart Automated Parking Gate System with Vehicle Counting",
        "objective": "Design and build an autonomous parking gate system that detects vehicles, opens and closes a gate automatically, counts available parking spaces, and denies entry when the lot is full.",
        "scope": "This project is intended for small apartment, office, or campus parking areas. It simulates an automated gate controller using vehicle detection at entry and exit points, a servo-controlled barrier, and a live parking count display.",
        "significance": "1) Parking management is a practical daily problem in urban areas. 2) The project is easy to demonstrate physically, which is useful during evaluation. 3) It has enough logic to be accepted as a proper robotics project without becoming too risky. 4) There are many online examples for Arduino parking barriers, servo gates, IR counters, and ultrasonic detection.",
        "micro": "Arduino Uno",
        "sensors": "HC-SR04 ultrasonic sensor for entry detection; IR sensor for exit detection; LDR for day/night mode",
        "actuators": "SG90 servo motor for barrier gate; I2C LCD 16x2; LEDs and buzzer for traffic/status indication",
        "body": "Small parking model/frame",
        "additional": "Push button, breadboard, jumper wires, resistors, power supply",
        "cost": "Total (Approx.): 1,785 BDT",
        "f1": ("Vehicle Detection and Gate Opening", "The system detects a vehicle at the entry point and opens the gate automatically if parking space is available.", "1) The ultrasonic sensor detects an approaching vehicle at the entry lane. 2) Arduino checks the current parking count. 3) If the parking lot is not full, the servo motor raises the barrier gate. 4) After the vehicle passes, the gate closes automatically. 5) The LCD shows GATE OPEN, GATE CLOSED, or LOT FULL."),
        "f2": ("Exit Detection and Space Count Update", "An IR sensor at the exit point detects vehicles leaving and updates the available parking count.", "1) The IR sensor detects a vehicle passing through the exit side. 2) Arduino decreases the occupied-space count when a vehicle exits. 3) The LCD continuously displays available spots. 4) If the lot was previously full, the system returns to entry-allowed mode. 5) A push button can be used for manual count reset during testing or maintenance."),
        "f3": ("Day/Night Status Lighting and Alert Indication", "The LDR detects ambient light and changes the gate-area lighting and status display behavior between day and night mode.", "1) The LDR reads ambient brightness. 2) In daylight mode, only standard status LEDs are active. 3) In night mode, extra white LEDs turn on to illuminate the gate area. 4) When the lot is full, the red LED and buzzer alert the driver. 5) The LCD alternates between parking count, gate state, and day/night status."),
    },
]

doc = Document(str(TEMPLATE))

# Header fields
for p in doc.paragraphs:
    t = p.text.strip()
    if t.startswith("Section"):
        p.text = "Section : 3"
    elif t.startswith("Group"):
        p.text = "Group : 5"
    elif t.startswith("Semester"):
        p.text = "Semester : Spring_2026"
    elif t.startswith("Submission Date"):
        p.text = "Submission Date:"
    elif t.startswith("Instructors"):
        p.text = "Instructors: RDW, SMYA"

# Member details lines in template are bullet lines starting with Name:
member_lines = [
    "Name: Md.Abir Hasan Rohan    ID: 22101294",
    "Name: Md.Saadat Rahman       ID: 22201101",
    "Name: Tazwar Saif Chowdhury  ID: 22201373",
    "Name: Muntasir Mamun Sakib   ID: 22299146",
    "Name: Rufaida Faruque        ID: 23101005",
]
name_slots = [i for i,p in enumerate(doc.paragraphs) if p.text.strip().startswith("Name:")]
for i, idx in enumerate(name_slots[:5]):
    doc.paragraphs[idx].text = member_lines[i]

# Helper: write content after label paragraph

def set_after_label(label, value):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == label and i + 1 < len(doc.paragraphs):
            doc.paragraphs[i + 1].text = value
            return True
    return False

# Fill each idea block in order of labels as present in template
idea_starts = [i for i,p in enumerate(doc.paragraphs) if p.text.strip() in ["Project Idea 1", "Project Idea 2", "Project Idea 3"]]
for idea_idx, start in enumerate(idea_starts[:3]):
    idea = ideas[idea_idx]
    end = idea_starts[idea_idx + 1] if idea_idx + 1 < len(idea_starts) else len(doc.paragraphs)
    block = doc.paragraphs[start:end]

    # local find inside block
    def find_idx(txt):
        for j, pp in enumerate(block):
            if pp.text.strip() == txt:
                return j
        return -1

    def set_next(txt, val):
        j = find_idx(txt)
        if j != -1 and j + 1 < len(block):
            block[j + 1].text = val

    set_next("1. Project Title:", idea["title"])
    set_next("Objective: [Describe the primary goal of the robot project]", f"Objective: {idea['objective']}")
    set_next("Scope: [Explain the intended use or application of the robot]", f"Scope: {idea['scope']}")
    set_next("Significance: [Discuss the importance and potential impact of the project]", f"Significance: {idea['significance']}")

    set_next("Microcontroller:", f"Microcontroller: {idea['micro']}")
    set_next("Sensors:", f"Sensors: {idea['sensors']}")
    set_next("Actuators:", f"Actuators: {idea['actuators']}")
    set_next("Body/Chassis:", f"Body/Chassis: {idea['body']}")
    set_next("Additional Components:", f"Additional Components: {idea['additional']}")

    set_next("4. Cost Breakdown (Approx.)", f"4. Cost Breakdown (Approx.): {idea['cost']}")

    # Functionality sections
    f1, f2, f3 = idea['f1'], idea['f2'], idea['f3']
    set_next("Functionality 1:", f"Functionality 1: {f1[0]}")
    set_next("Overview:", f"Overview: {f1[1]}")
    set_next("Working Procedure: [workflow, sensor integration, actuators logic, control logic, power management etc.]", f"Working Procedure: {f1[2]}")

    # second and third overview/workflow in block by occurrence order
    # find all exact matches
    overview_indices = [k for k,pp in enumerate(block) if pp.text.strip() == "Overview:"]
    wp_indices = [k for k,pp in enumerate(block) if pp.text.strip() == "Working Procedure:"]

    if len(overview_indices) >= 2 and overview_indices[1] + 1 < len(block):
        block[overview_indices[1] + 1].text = f"Overview: {f2[1]}"
    if len(wp_indices) >= 1 and wp_indices[0] + 1 < len(block):
        block[wp_indices[0] + 1].text = f"Working Procedure: {f2[2]}"

    set_next("Functionality 2:", f"Functionality 2: {f2[0]}")

    if len(overview_indices) >= 3 and overview_indices[2] + 1 < len(block):
        block[overview_indices[2] + 1].text = f"Overview: {f3[1]}"
    if len(wp_indices) >= 2 and wp_indices[1] + 1 < len(block):
        block[wp_indices[1] + 1].text = f"Working Procedure: {f3[2]}"

    set_next("Functionality 3:", f"Functionality 3: {f3[0]}")

# Save in place

doc.save(str(OUT))
print(f"Updated template: {OUT}")
