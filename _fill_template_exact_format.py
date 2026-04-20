from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt

SRC = Path(r"c:\Saadat\Class\Spring 2026\CSE461\Project\CSE461 Project Proposal Template.docx")
OUT = Path(r"c:\Saadat\Class\Spring 2026\CSE461\Project\CSE461 Project Proposal Template_FILLED_Group5.docx")

ideas = [
    {
        "title": "Smart Kitchen Safety and Ventilation System",
        "objective": "Design and build an autonomous kitchen safety system that uses multi-condition logic to distinguish normal cooking from dangerous events, reduce false alarms, and automatically activate ventilation, alarms, and gas-cutoff safety actions.",
        "scope": "The project is intended for household kitchens, dorm kitchens, or small food preparation areas. The system continuously monitors gas trend, temperature-humidity condition, and flame presence near the cooking area (with optional ambient-light context). Based on a multi-condition risk score and state-machine logic, it decides whether to open a vent, run an exhaust fan, trigger alarms, and execute gas cutoff.",
        "significance": "LPG gas leakage is a real and dangerous household problem in Bangladesh; most low-cost systems only raise an alarm, while this idea adds automatic action; it is socially useful and faculty-relevant; online implementation resources are widely available.",
        "micro": "Arduino Uno",
        "sensors": "MQ-2 gas/smoke sensor; DHT22 temperature and humidity sensor; Flame sensor module (IR); Optional LDR for kitchen activity / ambient-light context",
        "actuators": "SG90 servo motor for ventilation flap; Exhaust/DC fan controlled by relay; Solenoid gas valve for automatic cutoff (relay or MOSFET driver); I2C LCD 16x2",
        "body": "DIY enclosure/frame",
        "additional": "Stop Alarm push button (mandatory), LEDs, buzzer, relay module(s), power distribution board (recommended), resistors, breadboard, jumper wires, power supply",
        "cost_rows": [
            ("Controller", "700 BDT", "Arduino Uno"),
            ("Sensors", "650 BDT", "MQ-2, DHT22, flame sensor, optional LDR"),
            ("Actuators/Safety", "1000 BDT", "SG90, fan, solenoid valve, relay driver"),
            ("Display/Alert", "450 BDT", "LCD, LEDs, buzzer, stop button"),
            ("Board/Wires", "350 BDT", "Breadboard + jumpers + resistors + distribution board"),
            ("Frame/Power", "400 BDT", "Enclosure + power supply set"),
            ("Total", "3,550 BDT", "Approximate total"),
        ],
        "f1n": "Multi-Condition Hazard Detection with False Alarm Prevention",
        "f1o": "The system uses multiple sensors and state-based logic so normal cooking activity does not immediately trigger danger actions.",
        "f1w": "1) Arduino reads MQ-2, DHT22, and flame-sensor values continuously (optional LDR context can be added). 2) Sensor data is stabilized using warm-up lockout, moving-average filtering, and outlier rejection. 3) A weighted risk score determines states such as NORMAL, COOKING_NORMAL, WARNING, and ALARM_ACTIVE. 4) Persistence timers and hysteresis prevent one-sample spikes from creating false alarms. 5) Only sustained multi-condition danger transitions the system toward cutoff-level response.",
        "f2n": "Automatic Cutoff and Safety Actuation",
        "f2o": "When confirmed danger persists, the system performs automatic protective actions including gas cutoff.",
        "f2w": "1) If risk remains high for the configured confirmation window, the system enters ALARM_ACTIVE and then CUTOFF_LOCKED. 2) The solenoid gas valve is closed immediately via relay/driver. 3) The exhaust fan is turned on and vent flap is opened using the servo motor. 4) Buzzer, LEDs, and LCD provide clear alarm and hazard status. 5) Gas restore is blocked until a continuous safe window is observed and manual reset conditions are met.",
        "f3n": "Stop Alarm Button and Safe Recovery Workflow",
        "f3o": "The Stop Alarm button improves usability without allowing users to bypass safety protection.",
        "f3w": "1) A short press of the Stop Alarm button mutes the buzzer for a temporary window. 2) During mute, safety actions remain active: fan on, vent open, and gas cutoff locked. 3) LCD shows safety-active status (for example, ALARM MUTED - SAFETY ACTIVE). 4) A long press is accepted only after safe readings persist for the configured recovery window. 5) On valid recovery, the system returns to normal state and allows controlled reset of alarm status.",
    },
    {
        "title": "Smart Classroom Occupancy and Environment Controller",
        "objective": "Build an autonomous room monitoring and control system that counts people entering and leaving a room, monitors environmental conditions, and automatically controls ventilation and lighting to reduce energy waste.",
        "scope": "The project is designed for classrooms, offices, labs, or meeting rooms. It detects entry and exit movement, tracks occupancy count, and adjusts fan, vent, and lighting according to room usage and environmental readings.",
        "significance": "This solves visible real-world energy waste in empty rooms, is directly relevant to academic buildings, demonstrates multi-sensor automation logic, and has strong online support.",
        "micro": "Arduino Uno",
        "sensors": "IR sensor pair for entry/exit counting; DHT11 temperature and humidity sensor; LDR for ambient light detection",
        "actuators": "DC fan controlled by relay; SG90 servo motor for vent/window flap; I2C LCD 16x2",
        "body": "DIY frame/enclosure",
        "additional": "White LEDs for room light simulation, buzzer, push button, relay module, breadboard, jumper wires, power supply",
        "cost_rows": [
            ("Controller", "700 BDT", "Arduino Uno"),
            ("Sensors", "226 BDT", "IR pair, DHT11, LDR"),
            ("Actuators", "421 BDT", "Fan, SG90, relay"),
            ("Display/Alert", "184 BDT", "LCD, LEDs, buzzer, button"),
            ("Board/Wires", "190 BDT", "Breadboard + jumpers"),
            ("Frame/Power", "350 BDT", "Frame + 5V supply"),
            ("Total", "2,071 BDT", "Approximate total"),
        ],
        "f1n": "Bidirectional Occupancy Counting",
        "f1o": "Two IR sensors identify whether a person is entering or leaving the room based on trigger order.",
        "f1w": "1) Two IR sensors are placed at the doorway. 2) Outer then inner trigger sequence counts entry. 3) Inner then outer sequence counts exit. 4) The LCD shows live occupancy and maximum capacity. 5) If occupancy reaches a limit, the buzzer sounds and the LCD shows a warning.",
        "f2n": "Temperature-Controlled Ventilation",
        "f2o": "When occupied, the fan and vent are controlled according to temperature and humidity.",
        "f2w": "1) DHT11 monitors temperature and humidity continuously. 2) If the room is empty, the system keeps the fan off and the vent closed. 3) High temperature while occupied turns on the fan via relay. 4) The servo opens the vent partially or fully depending on threshold. 5) The LCD displays occupancy and ventilation status.",
        "f3n": "Occupancy-Based Smart Lighting",
        "f3o": "The LDR controls room lights based on ambient light only when the room is occupied.",
        "f3w": "1) The LDR provides ambient light level. 2) If the room is empty, lights remain off. 3) Occupied and low-light condition turns on the white LEDs. 4) Moderate light can keep partial lighting logic. 5) The LCD alternates occupancy, environment, and lighting status.",
    },
    {
        "title": "Smart Automated Parking Gate System with Vehicle Counting",
        "objective": "Design and build an autonomous parking gate system that detects vehicles, opens and closes a gate automatically, counts available parking spaces, and denies entry when the lot is full.",
        "scope": "This project is intended for small apartment, office, or campus parking areas. It simulates an automated gate controller using entry and exit detection, servo barrier control, and a live parking count display.",
        "significance": "Parking management is a practical urban problem; this project demonstrates clear automation with strong demo value and abundant online references.",
        "micro": "Arduino Uno",
        "sensors": "HC-SR04 ultrasonic sensor for entry detection; IR sensor for exit detection; LDR for day/night mode",
        "actuators": "SG90 servo motor for barrier gate; I2C LCD 16x2; LEDs and buzzer for traffic/status indication",
        "body": "Small parking model/frame",
        "additional": "Push button, breadboard, jumper wires, resistors, power supply",
        "cost_rows": [
            ("Controller", "700 BDT", "Arduino Uno"),
            ("Sensors", "211 BDT", "HC-SR04, IR, LDR"),
            ("Actuators", "150 BDT", "SG90 servo"),
            ("Display/Alert", "179 BDT", "LCD, LEDs, buzzer"),
            ("Board/Wires", "195 BDT", "Breadboard, jumpers, button"),
            ("Frame/Power", "350 BDT", "Parking model + supply"),
            ("Total", "1,785 BDT", "Approximate total"),
        ],
        "f1n": "Vehicle Detection and Gate Opening",
        "f1o": "The system detects an incoming vehicle and opens the gate automatically if space is available.",
        "f1w": "1) The ultrasonic sensor detects an approaching vehicle at the entry lane. 2) Arduino checks the current parking count. 3) If the parking lot is not full, the servo raises the barrier. 4) After the vehicle passes, the gate closes automatically. 5) The LCD shows gate and lot status.",
        "f2n": "Exit Detection and Space Count Update",
        "f2o": "An IR sensor detects vehicles leaving and updates available space count.",
        "f2w": "1) The IR sensor detects a vehicle leaving at the exit lane. 2) Arduino decreases the occupied count. 3) The LCD updates the available parking count. 4) Once a spot is free, entry is allowed again. 5) The push button supports manual reset during testing.",
        "f3n": "Day/Night Status Lighting and Alert Indication",
        "f3o": "The LDR changes lighting and status behavior between day and night conditions.",
        "f3w": "1) The LDR reads ambient brightness. 2) Day mode uses standard status indication. 3) Night mode enables extra white-light indication. 4) A full lot activates the red LED and buzzer. 5) The LCD alternates parking count, gate state, and day or night mode.",
    },
]


def insert_paragraph_after(paragraph, text=""):
    new_p = deepcopy(paragraph._p)
    for child in list(new_p):
        if child.tag.endswith('r'):
            new_p.remove(child)
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.paragraphs[-1]
    new_para._p = new_p
    new_para.text = text
    return new_para


def format_cost_row(paragraph, left, middle, right):
    paragraph.text = ""
    pf = paragraph.paragraph_format
    pf.left_indent = Inches(0.1)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.tab_stops.add_tab_stop(Inches(3.15))
    pf.tab_stops.add_tab_stop(Inches(4.65))
    run = paragraph.add_run(f"{left}\t{middle}\t{right}")
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)


doc = Document(str(SRC))

header_updates = {
    2: "CSE461",
    3: "Section : 3",
    4: "Group : 5",
    5: "Semester : Spring_2026",
    6: "Robot Project Proposal",
    8: "Member Details:",
    9: "Name: Md.Abir Hasan Rohan    ID: 22101294",
    10: "Name: Md.Saadat Rahman       ID: 22201101",
    11: "Name: Tazwar Saif Chowdhury  ID: 22201373",
    12: "Name: Muntasir Mamun Sakib   ID: 22299146",
    13: "Name: Rufaida Faruque        ID: 23101005",
    15: "Submission Date: March 11, 2026",
    16: "Instructors: RDW, SMYA",
}
for i, value in header_updates.items():
    doc.paragraphs[i].text = value

starts = [19, 45, 70]
for start, idea in zip(starts, ideas):
    mapping = {
        start + 0: f"Project Idea {starts.index(start) + 1}",
        start + 1: "1. Project Title:",
        start + 2: idea["title"],
        start + 3: "2. Purpose",
        start + 4: f"Objective: {idea['objective']}",
        start + 5: f"Scope: {idea['scope']}",
        start + 6: f"Significance: {idea['significance']}",
        start + 7: "3. Components",
        start + 8: f"Microcontroller: {idea['micro']}",
        start + 9: f"Sensors: {idea['sensors']}",
        start + 10: f"Actuators: {idea['actuators']}",
        start + 11: f"Body/Chassis: {idea['body']}",
        start + 12: f"Additional Components: {idea['additional']}",
        start + 13: "4. Cost Breakdown (Approx.)",
        start + 14: "5. Functionality Breakdown",
        start + 15: f"Functionality 1: {idea['f1n']}",
        start + 16: f"Overview: {idea['f1o']}",
        start + 17: f"Working Procedure: {idea['f1w']}",
        start + 18: f"Functionality 2: {idea['f2n']}",
        start + 19: f"Overview: {idea['f2o']}",
        start + 20: f"Working Procedure: {idea['f2w']}",
        start + 21: f"Functionality 3: {idea['f3n']}",
        start + 22: f"Overview: {idea['f3o']}",
        start + 23: f"Working Procedure: {idea['f3w']}",
    }
    for idx, value in mapping.items():
        doc.paragraphs[idx].text = value

# Insert cost rows from bottom to top so paragraph positions above remain valid.
for start, idea in reversed(list(zip(starts, ideas))):
    cost_heading = doc.paragraphs[start + 13]
    row = insert_paragraph_after(cost_heading, "")
    format_cost_row(row, "Components", "Unit Cost", "Total")
    previous = row
    for left, middle, right in idea["cost_rows"]:
        new_row = insert_paragraph_after(previous, "")
        format_cost_row(new_row, left, middle, right)
        previous = new_row

# Overwrite the separate filled copy with the exact-template version.
doc.save(str(OUT))
print(f"Saved exact-format filled copy: {OUT}")
